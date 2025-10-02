#!/usr/bin/env python3
"""
Legal Forms Field Extraction System

This script extracts placeholder fields from legal form templates (HTML, DOCX, TXT, PDF)
and stores them in a MySQL database with metadata management, deduplication, and
intelligent field detection.

Author: Legal Forms Extraction System
License: MIT
"""

import os
import re
import sys
import json
from pathlib import Path
from typing import List, Dict, Set, Tuple, Optional
from collections import defaultdict
import warnings

# Environment and configuration
from dotenv import load_dotenv
load_dotenv()

# Database
import mysql.connector
from mysql.connector import Error as MySQLError

# Text processing
from slugify import slugify
from rapidfuzz import fuzz

# HTML parsing (for HTML forms)
from bs4 import BeautifulSoup

# Document parsing
try:
    from docx import Document
    from docx.text.paragraph import Paragraph
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    warnings.warn("python-docx not available. DOCX extraction disabled.")

# PDF parsing
try:
    import pdfplumber
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    warnings.warn("pdfplumber not available. PDF extraction disabled.")

# Optional semantic embeddings for deduplication
ENABLE_SEMANTICS = os.getenv('ENABLE_SEMANTICS', 'True').lower() == 'true'
if ENABLE_SEMANTICS:
    try:
        from sentence_transformers import SentenceTransformer
        SEMANTIC_MODEL = SentenceTransformer('all-MiniLM-L6-v2')
        SEMANTICS_AVAILABLE = True
    except ImportError:
        SEMANTICS_AVAILABLE = False
        warnings.warn("sentence-transformers not available. Semantic deduplication disabled.")
else:
    SEMANTICS_AVAILABLE = False

# Configuration
DB_CONFIG = {
    'host': os.getenv('DB_HOST', 'localhost'),
    'user': os.getenv('DB_USER', 'root'),
    'password': os.getenv('DB_PASSWORD', ''),
    'database': os.getenv('DB_NAME', 'legal_forms'),
}

FORMS_DIR = os.getenv('FORMS_DIR', './')
ACCEPT_NON_ITALIC_DOCX = os.getenv('ACCEPT_NON_ITALIC_DOCX', 'False').lower() == 'true'
ACCEPT_BRACES_TXT = os.getenv('ACCEPT_BRACES_TXT', 'True').lower() == 'true'
ACCEPT_BRACES_PDF = os.getenv('ACCEPT_BRACES_PDF', 'True').lower() == 'true'
PDF_STRICT_MODE = os.getenv('PDF_STRICT_MODE', 'False').lower() == 'true'
FUZZY_THRESHOLD = int(os.getenv('FUZZY_THRESHOLD', '85'))
SEMANTIC_THRESHOLD = float(os.getenv('SEMANTIC_THRESHOLD', '0.85'))

# Placeholder patterns
BRACE_PATTERN = re.compile(r'\{([^}]+)\}')

# Paragraph group detection patterns
PARAGRAPH_GROUP_PATTERNS = [
    re.compile(r'\{[^}]*(CHOOSE|SELECT|PICK)\s+(ONE|APPROPRIATE|THE\s+FOLLOWING)[^}]*\}', re.IGNORECASE),
    re.compile(r'\{[^}]*(ONE\s+OF\s+THE\s+FOLLOWING)[^}]*\}', re.IGNORECASE),
]

# Enumeration patterns for paragraph options
ENUM_PATTERNS = [
    re.compile(r'^\s*\(([A-Z])\)\s*'),  # (A), (B), (C)
    re.compile(r'^\s*\((\d+)\)\s*'),    # (1), (2), (3)
    re.compile(r'^\s*([A-Z])\.\s*'),    # A., B., C.
    re.compile(r'^\s*(\d+)\.\s*'),      # 1., 2., 3.
]

# Canonical field name mappings (for better human readability)
CANONICAL_MAPPINGS = {
    'plaintiff': 'plaintiff_name',
    'defendant': 'defendant_name',
    'case number': 'case_number',
    'cause number': 'cause_number',
    'court': 'court_name',
    'county': 'county_name',
    'date': 'date_field',
    'time': 'time_field',
    'amount': 'amount_value',
    'attorney': 'attorney_name',
    'address': 'address_field',
    'phone': 'phone_number',
    'email': 'email_address',
    'name of': 'name',
    'style of the case': 'case_style',
}

# Data type inference patterns
DATA_TYPE_PATTERNS = {
    'date': [
        re.compile(r'\bdate\b', re.IGNORECASE),
        re.compile(r'\bday\b', re.IGNORECASE),
        re.compile(r'\bmonth\b', re.IGNORECASE),
        re.compile(r'\byear\b', re.IGNORECASE),
    ],
    'integer': [
        re.compile(r'\bnumber\b', re.IGNORECASE),
        re.compile(r'\bcount\b', re.IGNORECASE),
        re.compile(r'\bquantity\b', re.IGNORECASE),
    ],
    'decimal': [
        re.compile(r'\bpercent', re.IGNORECASE),
        re.compile(r'\brate\b', re.IGNORECASE),
    ],
    'currency': [
        re.compile(r'\bamount\b', re.IGNORECASE),
        re.compile(r'\bdollar', re.IGNORECASE),
        re.compile(r'\bprice\b', re.IGNORECASE),
        re.compile(r'\bcost\b', re.IGNORECASE),
        re.compile(r'\bfee\b', re.IGNORECASE),
        re.compile(r'\bpayment\b', re.IGNORECASE),
    ],
    'email': [
        re.compile(r'\bemail\b', re.IGNORECASE),
        re.compile(r'\be-mail\b', re.IGNORECASE),
    ],
    'phone': [
        re.compile(r'\bphone\b', re.IGNORECASE),
        re.compile(r'\btelephone\b', re.IGNORECASE),
        re.compile(r'\bfax\b', re.IGNORECASE),
    ],
    'boolean': [
        re.compile(r'\byes/no\b', re.IGNORECASE),
        re.compile(r'\btrue/false\b', re.IGNORECASE),
    ],
}


class DatabaseManager:
    """Manages database connections and schema operations."""
    
    def __init__(self, config: Dict):
        self.config = config
        self.connection = None
        
    def connect(self):
        """Establish database connection."""
        try:
            self.connection = mysql.connector.connect(**self.config)
            print(f"✓ Connected to MySQL database '{self.config['database']}'")
        except MySQLError as e:
            print(f"✗ Failed to connect to MySQL: {e}")
            sys.exit(1)
    
    def close(self):
        """Close database connection."""
        if self.connection and self.connection.is_connected():
            self.connection.close()
            print("✓ Database connection closed")
    
    def execute(self, query: str, params: Tuple = None, fetch: bool = False):
        """Execute a query and optionally fetch results."""
        cursor = self.connection.cursor(dictionary=True)
        try:
            cursor.execute(query, params or ())
            if fetch:
                result = cursor.fetchall()
                cursor.close()
                return result
            else:
                self.connection.commit()
                lastrowid = cursor.lastrowid
                cursor.close()
                return lastrowid
        except MySQLError as e:
            print(f"✗ Query error: {e}")
            print(f"  Query: {query[:100]}...")
            cursor.close()
            raise
    
    def create_schema(self):
        """Create database schema if it doesn't exist."""
        print("\n=== Creating Database Schema ===")
        
        # Forms table
        self.execute("""
            CREATE TABLE IF NOT EXISTS forms (
                id INT AUTO_INCREMENT PRIMARY KEY,
                filename VARCHAR(512) NOT NULL,
                file_path VARCHAR(1024) NOT NULL,
                file_type VARCHAR(50) NOT NULL,
                form_title VARCHAR(512),
                extraction_date DATETIME DEFAULT CURRENT_TIMESTAMP,
                last_updated DATETIME DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
                UNIQUE KEY unique_filename (filename)
            )
        """)
        print("✓ Created 'forms' table")
        
        # Fields table (canonical fields)
        self.execute("""
            CREATE TABLE IF NOT EXISTS fields (
                id INT AUTO_INCREMENT PRIMARY KEY,
                canonical_name VARCHAR(255) NOT NULL,
                display_name VARCHAR(512),
                data_type VARCHAR(50) DEFAULT 'text',
                tooltip TEXT,
                created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                UNIQUE KEY unique_canonical (canonical_name)
            )
        """)
        print("✓ Created 'fields' table")
        
        # Synonyms table (field variations)
        self.execute("""
            CREATE TABLE IF NOT EXISTS synonyms (
                id INT AUTO_INCREMENT PRIMARY KEY,
                field_id INT NOT NULL,
                raw_text VARCHAR(1024) NOT NULL,
                source_form_id INT,
                FOREIGN KEY (field_id) REFERENCES fields(id) ON DELETE CASCADE,
                FOREIGN KEY (source_form_id) REFERENCES forms(id) ON DELETE SET NULL,
                UNIQUE KEY unique_synonym (field_id, raw_text(255))
            )
        """)
        print("✓ Created 'synonyms' table")
        
        # Form-Field relationship (normalized many-to-many)
        self.execute("""
            CREATE TABLE IF NOT EXISTS form_fields (
                id INT AUTO_INCREMENT PRIMARY KEY,
                form_id INT NOT NULL,
                field_id INT NOT NULL,
                occurrence_count INT DEFAULT 1,
                FOREIGN KEY (form_id) REFERENCES forms(id) ON DELETE CASCADE,
                FOREIGN KEY (field_id) REFERENCES fields(id) ON DELETE CASCADE,
                UNIQUE KEY unique_form_field (form_id, field_id)
            )
        """)
        print("✓ Created 'form_fields' table")
        
        # Field options (for choice/option fields like plaintiff/defendant)
        self.execute("""
            CREATE TABLE IF NOT EXISTS field_options (
                id INT AUTO_INCREMENT PRIMARY KEY,
                field_id INT NOT NULL,
                option_value VARCHAR(512) NOT NULL,
                display_text VARCHAR(1024),
                FOREIGN KEY (field_id) REFERENCES fields(id) ON DELETE CASCADE,
                UNIQUE KEY unique_field_option (field_id, option_value(255))
            )
        """)
        print("✓ Created 'field_options' table")
        
        # Paragraph groups (choice groups like "CHOOSE ONE OF THE FOLLOWING")
        self.execute("""
            CREATE TABLE IF NOT EXISTS paragraph_groups (
                id INT AUTO_INCREMENT PRIMARY KEY,
                group_name VARCHAR(512),
                description TEXT,
                created_at DATETIME DEFAULT CURRENT_TIMESTAMP
            )
        """)
        print("✓ Created 'paragraph_groups' table")
        
        # Paragraph options (the enumerated choices within a group)
        self.execute("""
            CREATE TABLE IF NOT EXISTS paragraph_options (
                id INT AUTO_INCREMENT PRIMARY KEY,
                group_id INT NOT NULL,
                option_label VARCHAR(255),
                option_text TEXT,
                display_order INT,
                FOREIGN KEY (group_id) REFERENCES paragraph_groups(id) ON DELETE CASCADE
            )
        """)
        print("✓ Created 'paragraph_options' table")
        
        # Fields within paragraph options
        self.execute("""
            CREATE TABLE IF NOT EXISTS paragraph_option_fields (
                id INT AUTO_INCREMENT PRIMARY KEY,
                paragraph_option_id INT NOT NULL,
                field_id INT NOT NULL,
                FOREIGN KEY (paragraph_option_id) REFERENCES paragraph_options(id) ON DELETE CASCADE,
                FOREIGN KEY (field_id) REFERENCES fields(id) ON DELETE CASCADE,
                UNIQUE KEY unique_para_field (paragraph_option_id, field_id)
            )
        """)
        print("✓ Created 'paragraph_option_fields' table")
        
        # Form-Paragraph relationship
        self.execute("""
            CREATE TABLE IF NOT EXISTS form_paragraph_groups (
                id INT AUTO_INCREMENT PRIMARY KEY,
                form_id INT NOT NULL,
                group_id INT NOT NULL,
                FOREIGN KEY (form_id) REFERENCES forms(id) ON DELETE CASCADE,
                FOREIGN KEY (group_id) REFERENCES paragraph_groups(id) ON DELETE CASCADE,
                UNIQUE KEY unique_form_group (form_id, group_id)
            )
        """)
        print("✓ Created 'form_paragraph_groups' table")
        
        # Wide table for reporting (dynamic columns)
        self.execute("""
            CREATE TABLE IF NOT EXISTS form_field_values_wide (
                id INT AUTO_INCREMENT PRIMARY KEY,
                form_id INT NOT NULL,
                FOREIGN KEY (form_id) REFERENCES forms(id) ON DELETE CASCADE,
                UNIQUE KEY unique_form (form_id)
            )
        """)
        print("✓ Created 'form_field_values_wide' table")
        
        print("✓ Schema creation complete")


class FieldExtractor:
    """Extracts placeholder fields from various document formats."""
    
    def __init__(self, db: DatabaseManager):
        self.db = db
        self.field_cache = {}  # Cache for canonical field IDs
        
    def extract_from_html(self, file_path: str) -> Set[str]:
        """Extract placeholders from HTML files (existing form format)."""
        fields = set()
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            soup = BeautifulSoup(content, 'lxml')
            
            # Find all italic tags that might contain placeholders
            for italic_tag in soup.find_all('i'):
                text = italic_tag.get_text()
                # Check if the italic text is inside braces
                parent_text = italic_tag.parent.get_text() if italic_tag.parent else text
                matches = BRACE_PATTERN.findall(parent_text)
                for match in matches:
                    # Only include if the match contains the italic text
                    if text.strip() in match:
                        fields.add(match.strip())
            
            # Also find direct brace patterns in case some are not in italic tags
            matches = BRACE_PATTERN.findall(content)
            for match in matches:
                # Check if it looks like a placeholder (typically has descriptive text)
                if len(match.strip()) > 0:
                    fields.add(match.strip())
                    
        except Exception as e:
            print(f"  ⚠ Error extracting from HTML {file_path}: {e}")
        
        return fields
    
    def extract_from_docx(self, file_path: str) -> Set[str]:
        """Extract placeholders from DOCX files."""
        if not DOCX_AVAILABLE:
            return set()
        
        fields = set()
        try:
            doc = Document(file_path)
            
            for para in doc.paragraphs:
                text = para.text
                matches = BRACE_PATTERN.findall(text)
                
                for match in matches:
                    # Check if italicized (if required)
                    if not ACCEPT_NON_ITALIC_DOCX:
                        # Check if any run in the paragraph containing this match is italic
                        is_italic = any(
                            run.italic and match in run.text
                            for run in para.runs
                        )
                        if not is_italic:
                            continue
                    
                    fields.add(match.strip())
                    
        except Exception as e:
            print(f"  ⚠ Error extracting from DOCX {file_path}: {e}")
        
        return fields
    
    def extract_from_txt(self, file_path: str) -> Set[str]:
        """Extract placeholders from TXT files."""
        if not ACCEPT_BRACES_TXT:
            return set()
        
        fields = set()
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            matches = BRACE_PATTERN.findall(content)
            for match in matches:
                fields.add(match.strip())
                
        except Exception as e:
            print(f"  ⚠ Error extracting from TXT {file_path}: {e}")
        
        return fields
    
    def extract_from_pdf(self, file_path: str) -> Set[str]:
        """Extract placeholders from PDF files."""
        if not PDF_AVAILABLE or not ACCEPT_BRACES_PDF:
            return set()
        
        fields = set()
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        matches = BRACE_PATTERN.findall(text)
                        for match in matches:
                            fields.add(match.strip())
                            
        except Exception as e:
            print(f"  ⚠ Error extracting from PDF {file_path}: {e}")
        
        return fields
    
    def extract_from_file(self, file_path: str) -> Tuple[Set[str], str]:
        """Extract fields from any supported file type."""
        ext = Path(file_path).suffix.lower()
        
        if ext == '.html' or ext == '.htm':
            return self.extract_from_html(file_path), 'html'
        elif ext == '.docx':
            return self.extract_from_docx(file_path), 'docx'
        elif ext == '.txt':
            return self.extract_from_txt(file_path), 'txt'
        elif ext == '.pdf':
            return self.extract_from_pdf(file_path), 'pdf'
        else:
            return set(), 'unknown'
    
    def infer_data_type(self, field_text: str) -> str:
        """Infer data type from field text."""
        field_lower = field_text.lower()
        
        for dtype, patterns in DATA_TYPE_PATTERNS.items():
            for pattern in patterns:
                if pattern.search(field_lower):
                    return dtype
        
        return 'text'
    
    def generate_canonical_name(self, raw_text: str) -> str:
        """Generate a canonical field name from raw text."""
        # Normalize text
        text = raw_text.lower().strip()
        
        # Apply canonical mappings
        for key, value in CANONICAL_MAPPINGS.items():
            if key in text:
                text = text.replace(key, value)
        
        # Remove special characters and create slug
        canonical = slugify(text, separator='_')
        
        # Ensure it starts with a letter (for SQL compatibility)
        if canonical and not canonical[0].isalpha():
            canonical = 'field_' + canonical
        
        # Limit length
        if len(canonical) > 200:
            canonical = canonical[:200]
        
        return canonical or 'field_unknown'
    
    def generate_tooltip(self, raw_text: str) -> str:
        """Generate a human-friendly tooltip for a field."""
        # Clean up the text
        text = raw_text.strip()
        
        # Remove excessive whitespace
        text = ' '.join(text.split())
        
        # Capitalize first letter
        if text:
            text = text[0].upper() + text[1:]
        
        return text
    
    def detect_options(self, raw_text: str) -> List[str]:
        """Detect if field contains options (e.g., plaintiff/defendant)."""
        options = []
        
        # Look for slash-separated options
        if '/' in raw_text:
            parts = raw_text.split('/')
            if len(parts) <= 5:  # Reasonable number of options
                options = [p.strip() for p in parts if p.strip()]
        
        return options
    
    def find_or_create_field(self, raw_text: str, form_id: int) -> int:
        """Find existing field or create new one with deduplication."""
        # Check cache first
        cache_key = raw_text.lower().strip()
        if cache_key in self.field_cache:
            return self.field_cache[cache_key]
        
        # Generate canonical name
        canonical_name = self.generate_canonical_name(raw_text)
        
        # Check for existing field with fuzzy matching
        existing_fields = self.db.execute(
            "SELECT id, canonical_name FROM fields",
            fetch=True
        )
        
        best_match_id = None
        best_score = 0
        
        for field in existing_fields:
            # Fuzzy match on canonical names
            score = fuzz.ratio(canonical_name, field['canonical_name'])
            if score > best_score and score >= FUZZY_THRESHOLD:
                best_score = score
                best_match_id = field['id']
        
        # Semantic matching (if enabled)
        if SEMANTICS_AVAILABLE and not best_match_id:
            try:
                embeddings = SEMANTIC_MODEL.encode([canonical_name] + [f['canonical_name'] for f in existing_fields])
                from numpy import dot
                from numpy.linalg import norm
                
                query_emb = embeddings[0]
                for i, field_emb in enumerate(embeddings[1:]):
                    similarity = dot(query_emb, field_emb) / (norm(query_emb) * norm(field_emb))
                    if similarity > SEMANTIC_THRESHOLD:
                        best_match_id = existing_fields[i]['id']
                        break
            except Exception as e:
                print(f"  ⚠ Semantic matching error: {e}")
        
        # Use existing field or create new one
        if best_match_id:
            field_id = best_match_id
        else:
            # Create new field
            display_name = self.generate_tooltip(raw_text)
            data_type = self.infer_data_type(raw_text)
            tooltip = self.generate_tooltip(raw_text)
            
            field_id = self.db.execute(
                """
                INSERT INTO fields (canonical_name, display_name, data_type, tooltip)
                VALUES (%s, %s, %s, %s)
                ON DUPLICATE KEY UPDATE id=LAST_INSERT_ID(id)
                """,
                (canonical_name, display_name, data_type, tooltip)
            )
        
        # Add synonym
        try:
            self.db.execute(
                """
                INSERT IGNORE INTO synonyms (field_id, raw_text, source_form_id)
                VALUES (%s, %s, %s)
                """,
                (field_id, raw_text, form_id)
            )
        except Exception as e:
            print(f"  ⚠ Could not add synonym: {e}")
        
        # Detect and store options
        options = self.detect_options(raw_text)
        for option in options:
            try:
                self.db.execute(
                    """
                    INSERT IGNORE INTO field_options (field_id, option_value, display_text)
                    VALUES (%s, %s, %s)
                    """,
                    (field_id, option, option)
                )
            except Exception as e:
                print(f"  ⚠ Could not add option: {e}")
        
        # Update wide table schema if needed
        self._ensure_wide_column(canonical_name)
        
        # Cache result
        self.field_cache[cache_key] = field_id
        
        return field_id
    
    def _ensure_wide_column(self, canonical_name: str):
        """Ensure column exists in wide table."""
        try:
            # Check if column exists
            result = self.db.execute(
                f"""
                SELECT COUNT(*) as cnt
                FROM INFORMATION_SCHEMA.COLUMNS
                WHERE TABLE_SCHEMA = %s
                AND TABLE_NAME = 'form_field_values_wide'
                AND COLUMN_NAME = %s
                """,
                (self.db.config['database'], canonical_name),
                fetch=True
            )
            
            if result[0]['cnt'] == 0:
                # Add column
                self.db.execute(
                    f"ALTER TABLE form_field_values_wide ADD COLUMN `{canonical_name}` TEXT"
                )
                print(f"  ✓ Added wide table column: {canonical_name}")
        except Exception as e:
            print(f"  ⚠ Could not ensure wide column {canonical_name}: {e}")


class FormProcessor:
    """Processes forms and extracts metadata."""
    
    def __init__(self, db: DatabaseManager, extractor: FieldExtractor):
        self.db = db
        self.extractor = extractor
    
    def process_file(self, file_path: str):
        """Process a single form file."""
        filename = Path(file_path).name
        print(f"\n• Processing: {filename}")
        
        # Extract fields
        raw_fields, file_type = self.extractor.extract_from_file(file_path)
        
        if not raw_fields:
            print(f"  ℹ No fields found")
            return
        
        print(f"  ✓ Extracted {len(raw_fields)} raw placeholders")
        
        # Get or create form record
        form_id = self.db.execute(
            """
            INSERT INTO forms (filename, file_path, file_type, form_title)
            VALUES (%s, %s, %s, %s)
            ON DUPLICATE KEY UPDATE
                id=LAST_INSERT_ID(id),
                file_path=VALUES(file_path),
                file_type=VALUES(file_type),
                last_updated=CURRENT_TIMESTAMP
            """,
            (filename, file_path, file_type, filename)
        )
        
        # Process each field
        field_counts = defaultdict(int)
        for raw_field in raw_fields:
            try:
                field_id = self.extractor.find_or_create_field(raw_field, form_id)
                field_counts[field_id] += 1
            except Exception as e:
                print(f"  ⚠ Could not process field '{raw_field[:50]}...': {e}")
        
        # Update form-field relationships
        for field_id, count in field_counts.items():
            try:
                self.db.execute(
                    """
                    INSERT INTO form_fields (form_id, field_id, occurrence_count)
                    VALUES (%s, %s, %s)
                    ON DUPLICATE KEY UPDATE occurrence_count=VALUES(occurrence_count)
                    """,
                    (form_id, field_id, count)
                )
            except Exception as e:
                print(f"  ⚠ Could not link field to form: {e}")
        
        # Ensure wide table row exists
        try:
            self.db.execute(
                """
                INSERT IGNORE INTO form_field_values_wide (form_id)
                VALUES (%s)
                """,
                (form_id,)
            )
        except Exception as e:
            print(f"  ⚠ Could not create wide table row: {e}")
        
        print(f"  ✓ Stored {len(field_counts)} unique fields")
    
    def process_directory(self, directory: str):
        """Process all supported files in a directory."""
        print(f"\n=== Processing Forms in: {directory} ===")
        
        path = Path(directory)
        supported_extensions = ['.html', '.htm', '.docx', '.txt', '.pdf']
        
        files = []
        for ext in supported_extensions:
            files.extend(path.glob(f'*{ext}'))
        
        if not files:
            print("ℹ No supported files found")
            return
        
        print(f"Found {len(files)} files to process")
        
        for file_path in sorted(files):
            try:
                self.process_file(str(file_path))
            except Exception as e:
                print(f"✗ Error processing {file_path.name}: {e}")


def main():
    """Main entry point."""
    print("=" * 60)
    print("Legal Forms Field Extraction System")
    print("=" * 60)
    
    # Verify configuration
    print("\n=== Configuration ===")
    print(f"Database: {DB_CONFIG['database']}@{DB_CONFIG['host']}")
    print(f"Forms directory: {FORMS_DIR}")
    print(f"Semantic deduplication: {'Enabled' if SEMANTICS_AVAILABLE else 'Disabled'}")
    print(f"Accept non-italic DOCX: {ACCEPT_NON_ITALIC_DOCX}")
    
    # Initialize database
    db = DatabaseManager(DB_CONFIG)
    db.connect()
    
    try:
        # Create schema
        db.create_schema()
        
        # Initialize processors
        extractor = FieldExtractor(db)
        processor = FormProcessor(db, extractor)
        
        # Process forms
        processor.process_directory(FORMS_DIR)
        
        # Print summary
        print("\n=== Extraction Summary ===")
        
        form_count = db.execute("SELECT COUNT(*) as cnt FROM forms", fetch=True)[0]['cnt']
        field_count = db.execute("SELECT COUNT(*) as cnt FROM fields", fetch=True)[0]['cnt']
        synonym_count = db.execute("SELECT COUNT(*) as cnt FROM synonyms", fetch=True)[0]['cnt']
        
        print(f"✓ Total forms processed: {form_count}")
        print(f"✓ Total unique fields: {field_count}")
        print(f"✓ Total field synonyms: {synonym_count}")
        
        print("\n✓ Extraction complete!")
        print("\nNext steps:")
        print("1. Run admin UI: FLASK_APP=admin_app.py flask run --port 5000")
        print("2. Review and edit field tooltips in the admin interface")
        print("3. Query the database for reporting and analysis")
        
    except KeyboardInterrupt:
        print("\n\n⚠ Interrupted by user")
    except Exception as e:
        print(f"\n✗ Fatal error: {e}")
        import traceback
        traceback.print_exc()
    finally:
        db.close()


if __name__ == '__main__':
    main()
